"""
PROXVEL Backend - Import Service.
Parsea CSV subidos desde el dashboard → upsert en BD.
Soporta: absa, tags, clima, aforo.
Nuevos destinos detectados → is_active=False automáticamente.
"""
import csv
import io
import logging
import unicodedata
from contextvars import ContextVar
from typing import Literal

from sqlalchemy.orm import Session

from app.models.destination_model import Destination
from app.models.tourism_catalog_model import TourismCatalog
from app.models.aspect_score_model import DestinationAspectScore
from app.models.destination_tag_model import DestinationTag
from app.models.destination_month_context_model import DestinationMonthContext
from app.utils.normalizers import (
    generate_destination_id,
    destination_word_set,
    jaccard_similarity,
)
from app.services.destination_metric_sync_service import recalculate_status

logger = logging.getLogger(__name__)

ImportType = Literal["absa", "tags", "clima", "aforo", "catalog"]

# Selección del administrador al aplicar una importación (ver run_import):
#   _ALLOWED: destinos que se actualizan (None = todos los del archivo).
#   _CREATE_NEW: si se crean los destinos que aún no existen.
_ALLOWED: ContextVar[set[str] | None] = ContextVar("import_allowed", default=None)
_CREATE_NEW: ContextVar[bool] = ContextVar("import_create_new", default=True)


class SkipRow(Exception):
    """La fila es de un destino que el administrador no eligió actualizar."""


def _check_selected(destination_id: str, is_new: bool) -> None:
    allowed = _ALLOWED.get()
    if allowed is not None and destination_id not in allowed:
        raise SkipRow()
    if is_new and not _CREATE_NEW.get():
        raise SkipRow()

_ABSA_ASPECTS = [
    "accesibilidad", "aforo_multitudes", "alojamiento", "atencion_servicio",
    "atractivos", "clima", "costos", "gastronomia", "limpieza", "seguridad",
]

_MONTHS = {
    "enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
    "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
    "noviembre": 11, "diciembre": 12,
}


_SIMILARITY_THRESHOLD = 0.5


def _find_similar(all_dests: list[Destination], name: str) -> Destination | None:
    """
    Busca un destino existente cuyo nombre sea suficientemente similar al dado.
    Usa Jaccard sobre palabras significativas (umbral 0.5).
    Evita crear duplicados con nombres como "Kuélap" vs "Ciudadela de Kuélap".
    """
    query_words = destination_word_set(name)
    if not query_words:
        return None
    best_score = 0.0
    best_dest = None
    for dest in all_dests:
        score = jaccard_similarity(query_words, destination_word_set(dest.destination))
        if score > best_score:
            best_score = score
            best_dest = dest
    return best_dest if best_score >= _SIMILARITY_THRESHOLD else None


def _get_or_create_destination(
    db: Session,
    name: str,
    all_dests: list[Destination] | None = None,
) -> tuple[Destination, bool, str | None]:
    """
    Devuelve (destination, created, similar_name_warning).
    1. Busca por slug exacto.
    2. Busca por nombre exacto.
    3. Busca por similitud Jaccard ≥ 0.5 (previene duplicados).
    4. Si nada coincide, crea destino nuevo.
    """
    dest_id = generate_destination_id(name)
    dest = db.query(Destination).filter_by(destination_id=dest_id).first()
    if not dest:
        dest = db.query(Destination).filter_by(destination=name).first()
    if dest:
        _check_selected(dest.destination_id, False)
        return dest, False, None

    # Similitud: usa la lista pre-cargada si está disponible (evita N+1)
    candidates = all_dests if all_dests is not None else db.query(Destination).all()
    similar = _find_similar(candidates, name)
    if similar:
        warning = (
            f"'{name}' parece duplicado de '{similar.destination}' "
            f"(ID: {similar.destination_id}). Se usó el existente."
        )
        logger.warning("Import dedup: %s", warning)
        _check_selected(similar.destination_id, False)
        return similar, False, warning

    _check_selected(dest_id, True)

    dest = Destination(destination_id=dest_id, destination=name, is_active=False)
    db.add(dest)
    db.flush()
    db.add(TourismCatalog(destination_id=dest_id))
    candidates.append(dest)  # para que la misma fila no cree duplicado en el mismo CSV
    return dest, True, None


def _parse_csv(content: bytes) -> list[dict]:
    """Filas del archivo como diccionarios. Acepta CSV (UTF-8) o Excel .xlsx (primera hoja)."""
    if content[:2] == b"PK":  # .xlsx es un zip
        return _parse_xlsx(content)
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    return [row for row in reader]


def _parse_xlsx(content: bytes) -> list[dict]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as e:  # zip dañado u otro formato
        raise ValueError(f"no se pudo leer el Excel ({e.__class__.__name__}).")
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        return []
    keys = [str(h).strip() if h is not None else "" for h in header]
    out = []
    for values in rows:
        if values is None or all(v in (None, "") for v in values):
            continue
        row = {}
        for k, v in zip(keys, values):
            if not k:
                continue
            if isinstance(v, float) and v.is_integer() and k in ("mes_num", "month", "altitude_m"):
                v = int(v)
            row[k] = "" if v is None else str(v).strip()
        out.append(row)
    wb.close()
    return out


def _to_float(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


def _parse_month(row: dict) -> int | None:
    # "month" (formato anterior) o "mes_num" (CSV de la tesis, Fase 4).
    for key in ("month", "mes_num"):
        raw = row.get(key)
        if raw not in (None, ""):
            try:
                m = int(float(raw))
                return m if 1 <= m <= 12 else None
            except (ValueError, TypeError):
                pass
    name = (row.get("month_name") or row.get("mes") or "").strip().lower()
    return _MONTHS.get(name)


def _get_or_create_month_ctx(db: Session, dest_id: str, month: int) -> DestinationMonthContext:
    ctx = db.query(DestinationMonthContext).filter_by(
        destination_id=dest_id, month=month
    ).first()
    if not ctx:
        ctx = DestinationMonthContext(destination_id=dest_id, month=month)
        db.add(ctx)
    return ctx


def _import_absa_wide(db: Session, content: bytes) -> dict:
    rows = _parse_csv(content)
    all_dests = db.query(Destination).all()
    created = updated = errors = 0
    warnings: list[str] = []
    affected_ids: set[str] = set()

    for row in rows:
        name = (row.get("destination") or "").strip()
        if not name:
            errors += 1
            continue
        try:
            dest, is_new, warn = _get_or_create_destination(db, name, all_dests)
            if is_new:
                created += 1
            if warn:
                warnings.append(warn)
            scores = db.query(DestinationAspectScore).filter_by(
                destination_id=dest.destination_id
            ).first()
            if not scores:
                scores = DestinationAspectScore(destination_id=dest.destination_id)
                db.add(scores)
            for asp in _ABSA_ASPECTS:
                val = _to_float(row.get(asp))
                if val is not None:
                    setattr(scores, asp, val)
            affected_ids.add(dest.destination_id)
            updated += 1
        except SkipRow:
            continue
        except Exception as e:
            logger.error("import_absa row error: %s — %s", name, e)
            errors += 1

    db.flush()
    for did in affected_ids:
        recalculate_status(db, did)
    if affected_ids:
        db.flush()

    return {
        "type": "absa",
        "rows_processed": len(rows),
        "destinations_created": created,
        "scores_updated": updated,
        "errors": errors,
        "warnings": warnings,
    }


def _import_tags_legacy(db: Session, content: bytes) -> dict:
    """CSV: destination, tags_interes (string separado por comas). Reemplaza tags."""
    rows = _parse_csv(content)
    all_dests = db.query(Destination).all()
    created = updated = errors = 0
    warnings: list[str] = []

    for row in rows:
        name = (row.get("destination") or "").strip()
        tags_raw = (row.get("tags_interes") or "").strip()
        if not name:
            errors += 1
            continue
        try:
            dest, is_new, warn = _get_or_create_destination(db, name, all_dests)
            if is_new:
                created += 1
            if warn:
                warnings.append(warn)
            db.query(DestinationTag).filter_by(destination_id=dest.destination_id).delete()
            for tag in tags_raw.split(","):
                tag = tag.strip().lower()
                if tag:
                    db.add(DestinationTag(destination_id=dest.destination_id, tag=tag))
            updated += 1
        except SkipRow:
            continue
        except Exception as e:
            logger.error("import_tags row error: %s — %s", name, e)
            errors += 1
    db.flush()
    return {
        "type": "tags",
        "rows_processed": len(rows),
        "destinations_created": created,
        "records_updated": updated,
        "errors": errors,
        "warnings": warnings,
    }


# ── Contexto mensual (clima y aforo) ─────────────────────────────────────────
#
# Dos formatos:
#   * Tesis (Fase 4): una fila por destino × mes × preferencia (clima) o × tolerancia
#     (aforo), identificada por destination_id. Alimenta clima_por_preferencia y
#     aforo_por_tolerancia, que es lo que consume el motor.
#   * Anterior: una fila por destino × mes con un único score. Alimenta solo el score
#     escalar; se invalida el JSON de ese mes para que el motor no use datos viejos.

_CLIMATE_PREFS = {"frio", "templado", "calido", "seco", "indiferente"}
_TOLERANCES = {"baja", "media", "alta"}
_MAX_ERROR_SAMPLES = 10


def _norm_key(v) -> str:
    """Minúsculas y sin tildes: 'Cálido' -> 'calido'."""
    t = unicodedata.normalize("NFKD", str(v or "").strip().lower())
    return "".join(c for c in t if not unicodedata.combining(c))


def _score_in_range(v) -> float | None:
    f = _to_float(v)
    return f if f is not None and 0.0 <= f <= 1.0 else None


def _resolve_row_destination(
    db: Session, row: dict, all_dests: list[Destination], by_id: dict[str, Destination],
) -> tuple[Destination | None, bool, str | None, str | None]:
    """
    Devuelve (destino, creado, advertencia, error).
    Prioriza destination_id (CSV de la tesis); si no existe y hay nombre, lo crea con ese
    id (inactivo, fuera del motor). Sin id se resuelve por nombre (formato anterior).
    """
    did = (row.get("destination_id") or "").strip()
    name = (row.get("destination") or row.get("destino") or "").strip()

    if did:
        dest = by_id.get(did)
        if dest:
            _check_selected(did, False)
            return dest, False, None, None
        if not name:
            return None, False, None, f"El destino '{did}' no existe y la fila no trae nombre para crearlo."
        _check_selected(did, True)
        dest = Destination(destination_id=did, destination=name, is_active=False)
        db.add(dest)
        db.flush()
        db.add(TourismCatalog(destination_id=did))
        all_dests.append(dest)
        by_id[did] = dest
        return dest, True, None, None

    if not name:
        return None, False, None, "Fila sin destination_id ni nombre de destino."
    dest, created, warn = _get_or_create_destination(db, name, all_dests)
    by_id[dest.destination_id] = dest
    return dest, created, warn, None


def _import_context(db: Session, content: bytes, kind: str) -> dict:
    """kind: 'clima' | 'aforo'."""
    rows = _parse_csv(content)
    all_dests = db.query(Destination).all()
    by_id = {d.destination_id: d for d in all_dests}
    ctx_by: dict[tuple[str, int], DestinationMonthContext] = {
        (c.destination_id, c.month): c for c in db.query(DestinationMonthContext).all()
    }

    is_clima = kind == "clima"
    segment_col = "preferencia_climatica" if is_clima else "tolerancia_multitudes"
    valid_segments = _CLIMATE_PREFS if is_clima else _TOLERANCES
    json_attr = "clima_por_preferencia" if is_clima else "aforo_por_tolerancia"

    created = updated = errors = 0
    warnings: list[str] = []
    error_samples: list[str] = []
    touched: set[tuple[str, int]] = set()

    def fail(line: int, msg: str) -> None:
        nonlocal errors
        errors += 1
        if len(error_samples) < _MAX_ERROR_SAMPLES:
            error_samples.append(f"Fila {line}: {msg}")

    for line, row in enumerate(rows, start=2):  # la fila 1 es el encabezado
        month = _parse_month(row)
        if not month:
            fail(line, "mes inválido (usa mes_num 1-12 o el nombre del mes).")
            continue
        try:
            dest, is_new, warn, err = _resolve_row_destination(db, row, all_dests, by_id)
            if err:
                fail(line, err)
                continue
            if is_new:
                created += 1
            if warn:
                warnings.append(warn)

            did = dest.destination_id
            ctx = ctx_by.get((did, month))
            if not ctx:
                ctx = DestinationMonthContext(destination_id=did, month=month)
                db.add(ctx)
                ctx_by[(did, month)] = ctx

            segment = _norm_key(row.get(segment_col))
            if segment:  # ── formato de la tesis ──
                if segment not in valid_segments:
                    fail(line, f"{segment_col} '{segment}' no es válido (usa: {', '.join(sorted(valid_segments))}).")
                    continue
                score = _score_in_range(row.get("score_clima" if is_clima else "score_aforo_personalizado"))
                if score is None:
                    fail(line, "el puntaje debe ser un número entre 0 y 1.")
                    continue
                current = dict(getattr(ctx, json_attr) or {})
                current[segment] = score
                setattr(ctx, json_attr, current)  # reasignar: SQLAlchemy no detecta mutaciones in-place

                if is_clima:
                    if segment == "indiferente":
                        ctx.clima_score = score
                        ctx.clima_nivel = (row.get("nivel_clima") or "").strip() or ctx.clima_nivel
                        # Textos generales del mes (la fila «indiferente» no depende de la preferencia).
                        ctx.clima_explicacion = (row.get("explicacion_clima") or "").strip() or ctx.clima_explicacion
                        ctx.clima_riesgo = (row.get("riesgo_climatico") or "").strip() or ctx.clima_riesgo
                        tmax, tmin = _to_float(row.get("tmax_c")), _to_float(row.get("tmin_c"))
                        if tmax is not None and tmin is not None:
                            ctx.temp_prom = round((tmax + tmin) / 2, 1)
                        ctx.pp = _to_float(row.get("pp_mm")) if row.get("pp_mm") not in (None, "") else ctx.pp
                else:
                    base = _score_in_range(row.get("score_aforo"))
                    if base is not None:
                        ctx.aforo_score = base
                    ctx.aforo_nivel = (row.get("nivel_aforo") or "").strip() or ctx.aforo_nivel
                    ctx.aforo_fuente = "tesis_fase4"
                    ctx.aforo_explicacion = (row.get("explicacion_aforo") or "").strip() or ctx.aforo_explicacion
                    visitors = _to_float(row.get("visitantes_medios"))
                    if visitors is not None:
                        ctx.visitantes_medios = visitors
            else:  # ── formato anterior ──
                if is_clima:
                    raw = row.get("score_clima_modelo")
                    score = _score_in_range(raw if raw not in (None, "") else row.get("score_clima_base"))
                    if score is None:
                        fail(line, "falta score_clima_modelo/score_clima_base (0 a 1).")
                        continue
                    ctx.clima_score = score
                    ctx.clima_nivel = (row.get("nivel_clima_modelo") or row.get("nivel_clima") or "").strip() or None
                    ctx.temp_prom = _to_float(row.get("temp_prom"))
                    ctx.pp = _to_float(row.get("pp"))
                else:
                    raw = row.get("score_aforo_modelo")
                    score = _score_in_range(raw if raw not in (None, "") else row.get("score_aforo_base_promedio"))
                    if score is None:
                        fail(line, "falta score_aforo_modelo/score_aforo_base_promedio (0 a 1).")
                        continue
                    ctx.aforo_score = score
                    ctx.aforo_nivel = (row.get("nivel_aforo_modelo") or "").strip() or None
                    ctx.aforo_fuente = "oficial_visitantes"
                # El score escalar manda: se descarta el JSON de ese mes para no
                # dejar valores por perfil desactualizados.
                setattr(ctx, json_attr, None)

            touched.add((did, month))
            updated += 1
        except SkipRow:
            continue
        except Exception as e:
            logger.error("import_%s fila %s: %s", kind, line, e)
            fail(line, "no se pudo procesar la fila.")

    db.flush()

    # Aviso por destino: ¿quedó completo el contexto que usa el motor?
    expected = len(valid_segments)
    incomplete: dict[str, int] = {}
    for did, month in touched:
        data = getattr(ctx_by[(did, month)], json_attr) or {}
        if len(data) < expected:
            incomplete[did] = incomplete.get(did, 0) + 1
    for did, n in sorted(incomplete.items()):
        warnings.append(
            f"{did}: {n} mes(es) sin los {expected} valores de "
            f"{'preferencia climática' if is_clima else 'tolerancia a multitudes'}; "
            "el motor usará el puntaje general en esos meses."
        )

    return {
        "type": kind,
        "rows_processed": len(rows),
        "destinations_created": created,
        "records_updated": updated,
        "errors": errors,
        "error_samples": error_samples,
        "warnings": warnings,
    }


def import_clima(db: Session, content: bytes) -> dict:
    """CSV de clima → destination_month_context (clima_score y clima_por_preferencia)."""
    return _import_context(db, content, "clima")


def import_aforo(db: Session, content: bytes) -> dict:
    """CSV de aforo → destination_month_context (aforo_score y aforo_por_tolerancia)."""
    return _import_context(db, content, "aforo")


# ── ABSA y tags en el formato de la tesis ────────────────────────────────────

# Vocabulario de intereses de la tesis (Fase 3). Un tag fuera de esta lista se guarda,
# pero ningún perfil de la app lo declara, así que nunca suma coincidencias.
_INTEREST_VOCAB = {
    "academico", "aventura", "cultura", "familiar", "fotografia", "gastronomia",
    "historia", "montanismo", "naturaleza", "playa", "relax", "rural", "urbano",
}


def _headers(rows: list[dict]) -> set[str]:
    return set(rows[0].keys()) if rows else set()


def _import_absa_long(db: Session, content: bytes) -> dict:
    """
    matriz_destino_aspecto_sentimiento.csv (Fase 3): una fila por destino × aspecto.
    Usa score_normalizado_0_1 de TODAS las filas, como la tesis (no filtra por
    evidence_status); las celdas con evidencia insuficiente se informan como aviso.
    """
    rows = _parse_csv(content)
    all_dests = db.query(Destination).all()
    by_id = {d.destination_id: d for d in all_dests}
    created = updated = errors = 0
    warnings: list[str] = []
    error_samples: list[str] = []
    low_evidence: dict[str, list[str]] = {}
    scores_by_dest: dict[str, DestinationAspectScore] = {}
    affected: set[str] = set()

    def fail(line: int, msg: str) -> None:
        nonlocal errors
        errors += 1
        if len(error_samples) < _MAX_ERROR_SAMPLES:
            error_samples.append(f"Fila {line}: {msg}")

    if "score_normalizado_0_1" not in _headers(rows):
        raise ValueError("Falta la columna score_normalizado_0_1.")

    for line, row in enumerate(rows, start=2):
        aspect = _norm_key(row.get("aspecto"))
        if aspect not in _ABSA_ASPECTS:
            fail(line, f"aspecto '{aspect or '(vacío)'}' no es válido.")
            continue
        score = _score_in_range(row.get("score_normalizado_0_1"))
        if score is None:
            fail(line, "score_normalizado_0_1 debe ser un número entre 0 y 1.")
            continue
        try:
            dest, is_new, warn, err = _resolve_row_destination(db, row, all_dests, by_id)
            if err:
                fail(line, err)
                continue
            if is_new:
                created += 1
            if warn:
                warnings.append(warn)
            did = dest.destination_id
            rec = scores_by_dest.get(did)
            if rec is None:
                rec = db.query(DestinationAspectScore).filter_by(destination_id=did).first()
                if rec is None:
                    rec = DestinationAspectScore(destination_id=did)
                    db.add(rec)
                scores_by_dest[did] = rec
            setattr(rec, aspect, score)
            affected.add(did)
            updated += 1
            if _norm_key(row.get("evidence_status")) in ("evidencia_insuficiente", "baja_evidencia"):
                low_evidence.setdefault(did, []).append(aspect)
        except SkipRow:
            continue
        except Exception as e:
            logger.error("import_absa fila %s: %s", line, e)
            fail(line, "no se pudo procesar la fila.")

    db.flush()
    for did in affected:
        recalculate_status(db, did)
    if affected:
        db.flush()

    for did, aspects in sorted(low_evidence.items()):
        warnings.append(f"{did}: evidencia baja o insuficiente en {', '.join(sorted(set(aspects)))} (se cargó igual, como en la tesis).")

    return {
        "type": "absa",
        "rows_processed": len(rows),
        "destinations_created": created,
        "scores_updated": updated,
        "errors": errors,
        "error_samples": error_samples,
        "warnings": warnings,
    }


def _import_tags_registry(db: Session, content: bytes) -> dict:
    """
    registro_maestro_centros_turisticos.csv (Fase 3): destination_id, destination, intereses.
    `intereses` es una lista separada por comas. Reemplaza los tags del destino; una fila
    sin intereses se rechaza (no borra los tags que ya tiene).
    """
    rows = _parse_csv(content)
    all_dests = db.query(Destination).all()
    by_id = {d.destination_id: d for d in all_dests}
    created = updated = errors = 0
    warnings: list[str] = []
    error_samples: list[str] = []

    def fail(line: int, msg: str) -> None:
        nonlocal errors
        errors += 1
        if len(error_samples) < _MAX_ERROR_SAMPLES:
            error_samples.append(f"Fila {line}: {msg}")

    for line, row in enumerate(rows, start=2):
        tags = {t.strip().lower() for t in (row.get("intereses") or "").split(",") if t.strip()}
        if not tags:
            fail(line, "la fila no trae intereses.")
            continue
        try:
            dest, is_new, warn, err = _resolve_row_destination(db, row, all_dests, by_id)
            if err:
                fail(line, err)
                continue
            if is_new:
                created += 1
            if warn:
                warnings.append(warn)
            did = dest.destination_id
            db.query(DestinationTag).filter_by(destination_id=did).delete()
            for tag in sorted(tags):
                db.add(DestinationTag(destination_id=did, tag=tag))
            unknown = sorted(t for t in tags if _norm_key(t).replace(" ", "_") not in _INTEREST_VOCAB)
            if unknown:
                warnings.append(
                    f"{did}: tags fuera del vocabulario de la tesis ({', '.join(unknown)}); "
                    "ningún perfil de la app los declara, así que no sumarán coincidencias."
                )
            updated += 1
        except SkipRow:
            continue
        except Exception as e:
            logger.error("import_tags fila %s: %s", line, e)
            fail(line, "no se pudo procesar la fila.")

    db.flush()
    return {
        "type": "tags",
        "rows_processed": len(rows),
        "destinations_created": created,
        "records_updated": updated,
        "errors": errors,
        "error_samples": error_samples,
        "warnings": warnings,
    }


def import_absa(db: Session, content: bytes) -> dict:
    """ABSA: formato largo de la tesis (columna 'aspecto') o formato ancho anterior."""
    if "aspecto" in _headers(_parse_csv(content)):
        return _import_absa_long(db, content)
    return _import_absa_wide(db, content)


def import_tags(db: Session, content: bytes) -> dict:
    """Tags: registro maestro de la tesis (columna 'intereses') o formato anterior (tags_interes)."""
    headers = _headers(_parse_csv(content))
    if "intereses" in headers and "tags_interes" not in headers:
        return _import_tags_registry(db, content)
    return _import_tags_legacy(db, content)


# ── Ficha comercial de los destinos (Excel o CSV) ────────────────────────────

CATALOG_TEXT_FIELDS = (
    "city", "region", "category", "description", "activities_summary", "visitor_info_summary",
    "accessibility_summary", "recommended_duration", "how_to_get", "official_website",
)
CATALOG_COLUMNS = ("destination_id", "destination", *CATALOG_TEXT_FIELDS,
                   "altitude_m", "latitude", "longitude", "difficulty", "cost_level")
_CATEGORIES = {"arqueologia": "Arqueología", "naturaleza": "Naturaleza", "museo": "Museo", "recreacion": "Recreación"}
_DIFFICULTIES = {"facil", "moderada", "exigente"}
_COSTS = {"bajo", "medio", "alto"}


def _catalog_values(row: dict) -> dict:
    """Valida y convierte las celdas con valor de una fila de la ficha. Lanza ValueError."""
    values: dict = {}
    for f in CATALOG_TEXT_FIELDS:
        v = (row.get(f) or "").strip()
        if v:
            values[f] = v
    if "category" in values:
        cat = _CATEGORIES.get(_norm_key(values["category"]))
        if not cat:
            raise ValueError(f"categoría '{values['category']}' no válida (Arqueología, Naturaleza, Museo o Recreación).")
        values["category"] = cat
    if "official_website" in values and not values["official_website"].startswith("https://"):
        raise ValueError("official_website debe empezar con https://.")
    for f, lo, hi in (("latitude", -19.0, 0.5), ("longitude", -82.0, -68.0)):
        raw = (row.get(f) or "").strip()
        if raw:
            num = _to_float(raw)
            if num is None or not lo <= num <= hi:
                raise ValueError(f"{f} '{raw}' está fuera del Perú.")
            values[f] = num
    raw = (row.get("altitude_m") or "").strip()
    if raw:
        num = _to_float(raw)
        if num is None or not 0 <= num <= 7000:
            raise ValueError(f"altitude_m '{raw}' no es válida (0 a 7000).")
        values["altitude_m"] = int(round(num))
    for f, allowed in (("difficulty", _DIFFICULTIES), ("cost_level", _COSTS)):
        raw = _norm_key(row.get(f))
        if raw:
            if raw not in allowed:
                raise ValueError(f"{f} '{row.get(f)}' no es válido (usa {', '.join(sorted(allowed))}).")
            values[f] = raw
    return values


def _import_catalog(db: Session, content: bytes) -> dict:
    """
    Ficha comercial: una fila por destino, identificada por destination_id. Solo se escriben
    las celdas con valor (una celda vacía no borra lo que ya hay). Un destination_id nuevo
    con nombre crea el destino inactivo y fuera del motor.
    """
    rows = _parse_csv(content)
    if rows and "destination_id" not in _headers(rows):
        raise ValueError("Falta la columna destination_id.")
    all_dests = db.query(Destination).all()
    by_id = {d.destination_id: d for d in all_dests}
    created = updated = errors = 0
    warnings: list[str] = []
    error_samples: list[str] = []

    def fail(line: int, msg: str) -> None:
        nonlocal errors
        errors += 1
        if len(error_samples) < _MAX_ERROR_SAMPLES:
            error_samples.append(f"Fila {line}: {msg}")

    unknown_cols = sorted(c for c in _headers(rows) if c and c not in CATALOG_COLUMNS)
    if unknown_cols:
        warnings.append(f"Columnas ignoradas (no son de la ficha): {', '.join(unknown_cols)}.")

    for line, row in enumerate(rows, start=2):
        if not (row.get("destination_id") or "").strip():
            fail(line, "falta destination_id.")
            continue
        try:
            values = _catalog_values(row)
            dest, is_new, warn, err = _resolve_row_destination(db, row, all_dests, by_id)
            if err:
                fail(line, err)
                continue
            if is_new:
                created += 1
            if warn:
                warnings.append(warn)
            db.flush()  # el destino recién creado ya trae su ficha vacía pendiente de guardar
            cat = db.query(TourismCatalog).filter_by(destination_id=dest.destination_id).first()
            if cat is None:
                cat = TourismCatalog(destination_id=dest.destination_id)
                db.add(cat)
            name = (row.get("destination") or "").strip()
            if name and not is_new:
                dest.destination = name
            for f, v in values.items():
                setattr(cat, f, v)
            updated += 1
        except SkipRow:
            continue
        except ValueError as e:
            fail(line, str(e))
        except Exception as e:
            logger.error("import_catalog fila %s: %s", line, e)
            fail(line, "no se pudo procesar la fila.")

    db.flush()
    return {
        "type": "catalog",
        "rows_processed": len(rows),
        "destinations_created": created,
        "records_updated": updated,
        "errors": errors,
        "error_samples": error_samples,
        "warnings": warnings,
    }


def export_catalog_rows(db: Session) -> list[dict]:
    """Ficha comercial actual en el formato del importador (plantilla con los datos de hoy)."""
    cats = {c.destination_id: c for c in db.query(TourismCatalog).all()}
    out = []
    for d in db.query(Destination).order_by(Destination.destination_id).all():
        c = cats.get(d.destination_id)
        row = {"destination_id": d.destination_id, "destination": d.destination}
        for f in CATALOG_COLUMNS[2:]:
            v = getattr(c, f, None) if c else None
            row[f] = "" if v is None else v
        out.append(row)
    return out


# ── Evaluación previa y aplicación selectiva ─────────────────────────────────

def _snapshot(db: Session, kind: str) -> dict[str, object]:
    """Estado comparable, por destino, de los datos que toca cada importador."""
    snap: dict[str, object] = {}
    if kind == "absa":
        for r in db.query(DestinationAspectScore).all():
            snap[r.destination_id] = tuple(round(getattr(r, a) or -1, 6) for a in _ABSA_ASPECTS)
    elif kind == "tags":
        tags: dict[str, list[str]] = {}
        for t in db.query(DestinationTag).all():
            tags.setdefault(t.destination_id, []).append(t.tag)
        snap = {k: tuple(sorted(v)) for k, v in tags.items()}
    elif kind in ("clima", "aforo"):
        cols = (("clima_score", "clima_nivel", "clima_por_preferencia", "temp_prom", "pp")
                if kind == "clima" else
                ("aforo_score", "aforo_nivel", "aforo_por_tolerancia", "visitantes_medios"))
        per: dict[str, list] = {}
        for c in db.query(DestinationMonthContext).all():
            vals = [repr(sorted(v.items())) if isinstance(v, dict) else v for v in (getattr(c, k) for k in cols)]
            per.setdefault(c.destination_id, []).append((c.month, *vals))
        snap = {k: tuple(sorted(v)) for k, v in per.items()}
    elif kind == "catalog":
        names = dict(db.query(Destination.destination_id, Destination.destination).all())
        for c in db.query(TourismCatalog).all():
            snap[c.destination_id] = (names.get(c.destination_id), *[getattr(c, f, None) for f in CATALOG_COLUMNS[2:]])
    return snap


def _file_destination_ids(content: bytes) -> set[str]:
    """destination_id que trae el archivo (para listar también los que no cambian)."""
    try:
        return {(r.get("destination_id") or "").strip() for r in _parse_csv(content)} - {""}
    except Exception:
        return set()


_IMPORTERS = {
    "absa": lambda db, c: import_absa(db, c),
    "tags": lambda db, c: import_tags(db, c),
    "clima": lambda db, c: import_clima(db, c),
    "aforo": lambda db, c: import_aforo(db, c),
    "catalog": lambda db, c: _import_catalog(db, c),
}


def run_import(
    db: Session,
    import_type: str,
    content: bytes,
    *,
    apply: bool,
    destinations: set[str] | None = None,
    create_new: bool = True,
) -> dict:
    """
    Evalúa (apply=False) o aplica (apply=True) un archivo.
    - Evaluación: corre la importación completa, compara antes/después por destino y deshace todo.
    - Aplicación: escribe solo los destinos elegidos (None = todos) y los datos del archivo
      reemplazan a los existentes (actualizar). create_new decide si se crean destinos nuevos.
    """
    importer = _IMPORTERS.get(import_type)
    if importer is None:
        raise ValueError(f"Tipo no soportado: {import_type}")

    before = _snapshot(db, import_type)
    known = {d.destination_id for d in db.query(Destination).all()}
    t_allowed = _ALLOWED.set(destinations if apply else None)
    t_new = _CREATE_NEW.set(create_new if apply else True)
    try:
        result = importer(db, content)
        db.flush()
        after = _snapshot(db, import_type)
        names = dict(db.query(Destination.destination_id, Destination.destination).all())
        changed = {k for k in set(after) | set(before) if after.get(k) != before.get(k)}
        listed = changed | (_file_destination_ids(content) & set(names))
        detail = []
        for did in sorted(listed, key=lambda k: (names.get(k) or k).lower()):
            status = "nuevo" if did not in known else ("cambia" if did in changed else "sin_cambios")
            detail.append({"destination_id": did, "destination": names.get(did, did), "status": status})
        result["destinations"] = detail
        result["summary"] = {
            "nuevos": sum(1 for d in detail if d["status"] == "nuevo"),
            "cambian": sum(1 for d in detail if d["status"] == "cambia"),
            "sin_cambios": sum(1 for d in detail if d["status"] == "sin_cambios"),
        }
        result["mode"] = "apply" if apply else "preview"
        if apply:
            db.commit()
        else:
            db.rollback()
        return result
    except Exception:
        db.rollback()
        raise
    finally:
        _ALLOWED.reset(t_allowed)
        _CREATE_NEW.reset(t_new)


def import_csv(db: Session, import_type: ImportType, content: bytes) -> dict:
    """Aplica el archivo completo de una vez (todos los destinos)."""
    return run_import(db, import_type, content, apply=True)

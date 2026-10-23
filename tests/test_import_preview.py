"""
Importación en dos pasos desde el dashboard: evaluar el archivo (sin guardar) y aplicarlo a
todos o a algunos destinos. Incluye la ficha comercial en Excel. Restaura lo que modifica.
Necesita los datos de `python scripts/seed_from_thesis.py`; si la BD está vacía se omite.
"""
import io
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.main import app
from app.config.database import SessionLocal
from app.core.security import create_access_token
from app.models.destination_model import Destination
from app.models.destination_month_context_model import DestinationMonthContext
from app.models.tourism_catalog_model import TourismCatalog
from app.models.user_model import User
from app.utils.hash_util import get_password_hash

API = "/api/v1"
ADMIN = "ZPRV_ADMIN"
A, B = "huaca_luna", "chan_chan"
SEED = Path(__file__).resolve().parent.parent / "seed_data"


@pytest.fixture(scope="module")
def client():
    with TestClient(app) as c:
        yield c


def _ficha(db, did):
    c = db.query(TourismCatalog).filter_by(destination_id=did).first()
    return {"recommended_duration": c.recommended_duration, "cost_level": c.cost_level,
            "altitude_m": c.altitude_m, "category": c.category}


@pytest.fixture(scope="module", autouse=True)
def entorno():
    db = SessionLocal()
    if db.query(Destination).filter_by(destination_id=A).first() is None:
        db.close()
        pytest.skip("BD sin datos: ejecuta scripts/seed_from_thesis.py")
    db.query(User).filter_by(user_id=ADMIN).delete()
    db.add(User(user_id=ADMIN, name=ADMIN, email="zprv.admin@mailinator.com",
                password_hash=get_password_hash("ClaveSegura123"), role="admin",
                is_active=True, account_status="active"))
    db.commit()
    originales = {d: _ficha(db, d) for d in (A, B)}
    ctx_original = {
        (c.destination_id, c.month): (c.aforo_por_tolerancia, c.aforo_score, c.aforo_nivel, c.visitantes_medios)
        for c in db.query(DestinationMonthContext).filter(DestinationMonthContext.destination_id.in_([A, B])).all()
    }
    yield
    db.expire_all()  # la API cambió datos en otras sesiones: releer antes de restaurar
    for did, vals in originales.items():
        cat = db.query(TourismCatalog).filter_by(destination_id=did).first()
        for k, v in vals.items():
            setattr(cat, k, v)
    for (did, m), (tol, score, nivel, vis) in ctx_original.items():
        c = db.query(DestinationMonthContext).filter_by(destination_id=did, month=m).first()
        c.aforo_por_tolerancia, c.aforo_score, c.aforo_nivel, c.visitantes_medios = tol, score, nivel, vis
    db.query(TourismCatalog).filter_by(destination_id="zprv_nuevo").delete()
    db.query(Destination).filter_by(destination_id="zprv_nuevo").delete()
    db.query(User).filter_by(user_id=ADMIN).delete()
    db.commit()
    db.close()


def h():
    return {"Authorization": f"Bearer {create_access_token(data={'sub': ADMIN, 'role': 'admin'})}"}


def xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


HEADER = ["destination_id", "destination", "recommended_duration", "cost_level", "altitude_m", "category"]


def ficha_excel():
    return xlsx([
        HEADER,
        [A, "", "90 minutos", "medio", 65, "arqueologia"],
        [B, "", "3 horas", "alto", "", ""],
        ["zprv_nuevo", "Destino de prueba", "1 hora", "bajo", 100, "Museo"],
    ])


def post(client, path, content, name="ficha.xlsx", **form):
    return client.post(f"{API}/admin/import/{path}", files={"file": (name, content)}, data=form, headers=h())


def test_evaluar_no_guarda_y_dice_que_cambia(client):
    with SessionLocal() as db:
        antes = _ficha(db, A)
    r = post(client, "catalog/preview", ficha_excel())
    assert r.status_code == 200, r.text
    data = r.json()["data"]
    assert data["mode"] == "preview"
    estados = {d["destination_id"]: d["status"] for d in data["destinations"]}
    assert estados == {A: "cambia", B: "cambia", "zprv_nuevo": "nuevo"}
    assert data["summary"] == {"nuevos": 1, "cambian": 2, "sin_cambios": 0}
    with SessionLocal() as db:
        assert _ficha(db, A) == antes, "la evaluación no debe escribir nada"
        assert db.query(Destination).filter_by(destination_id="zprv_nuevo").first() is None


def test_aplicar_solo_a_los_elegidos_y_sin_crear(client):
    with SessionLocal() as db:
        antes_b = _ficha(db, B)
    r = post(client, "catalog", ficha_excel(), destinations=A, create_new="false")
    assert r.status_code == 200, r.text
    with SessionLocal() as db:
        a = _ficha(db, A)
        assert (a["recommended_duration"], a["cost_level"], a["altitude_m"], a["category"]) == \
            ("90 minutos", "medio", 65, "Arqueología")
        assert _ficha(db, B) == antes_b, "un destino no elegido no se toca"
        assert db.query(Destination).filter_by(destination_id="zprv_nuevo").first() is None


def test_aplicar_todo_sobrescribe_y_crea_inactivo(client):
    r = post(client, "catalog", ficha_excel())
    assert r.status_code == 200
    with SessionLocal() as db:
        assert _ficha(db, B)["recommended_duration"] == "3 horas"
        assert _ficha(db, B)["altitude_m"] is not None, "celda vacía no borra lo que había"
        nuevo = db.query(Destination).filter_by(destination_id="zprv_nuevo").first()
        assert nuevo and not nuevo.is_active and not nuevo.in_engine
    # Volver a evaluar el mismo archivo: ya no hay cambios.
    data = post(client, "catalog/preview", ficha_excel()).json()["data"]
    assert data["summary"]["cambian"] == 0 and data["summary"]["nuevos"] == 0


def test_ficha_valida_celdas(client):
    malo = xlsx([HEADER + ["latitude"], [A, "", "", "gratis", "", "playa", "40.0"]])
    data = post(client, "catalog/preview", malo).json()["data"]
    assert data["errors"] == 1 and data["error_samples"]


def test_plantilla_con_datos_actuales(client):
    r = client.get(f"{API}/admin/import/catalog/template", headers=h())
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    filas = list(ws.iter_rows(values_only=True))
    assert filas[0][0] == "destination_id"
    ids = {f[0] for f in filas[1:]}
    assert {A, B} <= ids
    # La plantilla descargada se puede volver a subir tal cual: no cambia nada.
    data = post(client, "catalog/preview", r.content).json()["data"]
    assert data["errors"] == 0 and data["summary"]["cambian"] == 0


def test_snapshot_aforo_es_solo_lectura(client):
    with SessionLocal() as db:
        before = dict(db.query(DestinationMonthContext).filter_by(destination_id=A,month=1).first().aforo_por_tolerancia)
    contenido = (SEED / "aforo_scores.csv").read_bytes()
    assert post(client, "aforo/preview", contenido, name="aforo_scores.csv").status_code == 409
    assert post(client, "aforo", contenido, name="aforo_scores.csv", destinations=A).status_code == 409
    with SessionLocal() as db:
        assert db.query(DestinationMonthContext).filter_by(destination_id=A,month=1).first().aforo_por_tolerancia == before
